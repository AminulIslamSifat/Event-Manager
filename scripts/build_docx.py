from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = '/home/sifat/sable_output/assignment/assignment.docx'
ASSETS = '/home/sifat/sable_output/assignment/assets'

DARK = RGBColor(0x0d, 0x11, 0x17)
BLUE = RGBColor(0x1f, 0x6f, 0xeb)
BLUE_DK = RGBColor(0x0b, 0x4e, 0xa2)
GREY = RGBColor(0x6a, 0x73, 0x7d)
WHITE = RGBColor(0xff, 0xff, 0xff)

HEAD_FILL = '1F6FEB'
ALT_FILL = 'F7F9FC'
CALLOUT_FILL = 'F0F6FF'


def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hexfill)
    tcPr.append(sh)


def cell_borders(cell, color='D8DEE6', sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)


def para_shade(p, hexfill, left_bar=None):
    pPr = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), hexfill)
    pPr.append(sh)
    if left_bar:
        pbdr = OxmlElement('w:pBdr')
        lb = OxmlElement('w:left')
        lb.set(qn('w:val'), 'single')
        lb.set(qn('w:sz'), '18')
        lb.set(qn('w:space'), '8')
        lb.set(qn('w:color'), left_bar)
        pbdr.append(lb)
        pPr.append(pbdr)


def build():
    doc = Document()

    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    def add(text, size=10.5, bold=False, italic=False, color=DARK,
            align=None, before=0, after=6, bar=None, fill=None, indent=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        if fill or bar:
            para_shade(p, fill or 'FFFFFF', bar)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color
        return p

    def table(rows, widths=None, header=True):
        t = doc.add_table(rows=0, cols=len(rows[0]))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci, val in enumerate(row):
                c = cells[ci]
                c.text = ''
                p = c.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(str(val))
                r.font.size = Pt(9.6)
                if header and ri == 0:
                    r.bold = True
                    r.font.color.rgb = WHITE
                    shade(c, HEAD_FILL)
                    cell_borders(c, '1A5EC4')
                else:
                    r.font.color.rgb = DARK
                    if ci == 0:
                        r.bold = True
                    shade(c, ALT_FILL if ri % 2 == 0 else 'FFFFFF')
                    cell_borders(c)
        if widths:
            for ri in range(len(t.rows)):
                for ci, w in enumerate(widths):
                    t.rows[ri].cells[ci].width = Cm(w)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def image(fname, wcm, caption):
        path = os.path.join(ASSETS, fname)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(path, width=Cm(wcm))
        add(caption, size=8.5, italic=True, color=GREY,
            align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    # ---------------- title ----------------
    add('Hometown Map \u2014 Graph Assignment', size=23, bold=True, color=DARK,
        align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    add('CSE Discrete Mathematics', size=13, bold=True, color=BLUE,
        align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add('Sifat  |  October 2026', size=9.5, color=GREY,
        align=WD_ALIGN_PARAGRAPH.CENTER, after=16)

    # ---------------- Q1 ----------------
    add('Q1 \u2014 Map, Checkpoints & Road Distances', size=16, bold=True,
        color=DARK, before=4, after=8, bar=HEAD_FILL, fill='EEF4FF', indent=0.3)
    add('1.1 Annotated Map', size=12.5, bold=True, color=BLUE, before=8, after=4)
    add('Ten checkpoints (A\u2013J) marked on the map, with G (my house) positioned in the centre. Roads connecting the checkpoints are traced, and each road is labelled with its length.')
    image('01_map.png', 13.5, 'Annotated map with 10 checkpoints, road traces, and distance labels')

    add('1.2 Checkpoint Reference', size=12.5, bold=True, color=BLUE, before=8, after=4)
    table([
        ['Node', 'Location', 'Type'],
        ['A', 'Thana Rd (west)', 'Junction'],
        ['B', 'Thana Rd x Kazi Nazrul Islam Rd', 'Junction'],
        ['C', 'Farhan Decoration (south spur)', 'Road point'],
        ['D', 'Thana Rd (far west)', 'Road point'],
        ['E', 'Dynamic School & College', 'Landmark'],
        ['F', 'Abu Bakr Siddik Jame Masjid', 'Landmark'],
        ['G', 'My House (centre node)', 'Centre'],
        ['H', 'Hazipara Water Pump', 'Junction'],
        ['I', 'Hazipara Al-Aqsa Jame Masjid', 'Landmark'],
        ['J', 'She & Space (clothing store)', 'Landmark'],
    ], widths=[2.2, 8.0, 4.0])

    add('1.3 Road Lengths', size=12.5, bold=True, color=BLUE, before=8, after=4)
    add('Lengths measured along the road path from the map (scale: 0.2733 m/pixel at zoom 19, latitude 23.788).')
    table([
        ['Road', 'Length (px)', 'Length (km)'],
        ['A\u2013B', '134', '0.037'], ['A\u2013C', '368', '0.101'],
        ['A\u2013E', '219', '0.060'], ['A\u2013H', '315', '0.086'],
        ['A\u2013D', '394', '0.108'], ['B\u2013F', '405', '0.111'],
        ['B\u2013D', '310', '0.085'], ['E\u2013G', '356', '0.097'],
        ['F\u2013G', '482', '0.132'], ['G\u2013J', '372', '0.102'],
        ['H\u2013I', '268', '0.073'], ['I\u2013J', '283', '0.077'],
    ], widths=[4.0, 5.0, 5.0])

    # ---------------- Q2 ----------------
    add('Q2 \u2014 Graph Model', size=16, bold=True, color=DARK,
        before=16, after=8, bar=HEAD_FILL, fill='EEF4FF', indent=0.3)
    add('2.1 Selected Model', size=12.5, bold=True, color=BLUE, before=8, after=4)
    add('My given graph is an undirected multigraph.', size=11.5, bold=True,
        color=BLUE_DK, fill=CALLOUT_FILL, bar=HEAD_FILL, indent=0.3, after=8)
    add('The graph that is not directed and has parallel edges is called an undirected multigraph. In my given graph, the edges are not directed and there are parallel edges between H and I, so it is an undirected multigraph.')

    add('2.2 Why Each Other Model Is Less Suitable', size=12.5, bold=True, color=BLUE, before=8, after=4)
    table([
        ['Model', 'Verdict', 'Reason'],
        ['Undirected simple graph', 'Fails', 'Cannot represent the parallel edges between H and I'],
        ['Undirected pseudograph', 'Overkill', 'Permits self-loops, but no road starts and ends at the same checkpoint'],
        ['Directed simple graph', 'Fails', 'Roads are two-way, so direction is meaningless \u2014 and parallel edges are still banned'],
        ['Directed multigraph', 'Fails', 'Parallel edges allowed, but direction is still meaningless for two-way roads'],
        ['Undirected multigraph', 'Selected', 'Parallel edges allowed, direction ignored \u2014 matches the given graph exactly'],
    ], widths=[4.3, 2.2, 8.0])

    add('2.3 The Graph', size=12.5, bold=True, color=BLUE, before=8, after=4)
    image('02_graph.png', 14.5, 'Undirected multigraph with 10 nodes and 13 edges. The parallel pair H\u2013I is drawn as two separate arcs.')
    table([
        ['Property', 'Value'],
        ['Nodes (|V|)', '10'],
        ['Edges (|E|)', '13'],
        ['Degree sum', '26 = 2 x 13'],
    ], widths=[6.0, 6.0])
    add('Degrees: A = 5, B = 3, C = 1, D = 2, E = 2, F = 2, G = 3, H = 3, I = 3, J = 2')

    # ---------------- Q3 ----------------
    add('Q3 \u2014 Cafe Assignment', size=16, bold=True, color=DARK,
        before=16, after=8, bar=HEAD_FILL, fill='EEF4FF', indent=0.3)
    add('3.1 Part A \u2014 Is the Assignment Possible?', size=12.5, bold=True, color=BLUE, before=8, after=4)
    add('No \u2014 the assignment is impossible.', size=11.5, bold=True,
        color=BLUE_DK, fill=CALLOUT_FILL, bar=HEAD_FILL, indent=0.3, after=8)
    add("No, the cafe can't be assigned the given way for my hometown graph. Propagating a 2-colouring from A forces B and D to the same colour, but the road B\u2013D joins them. The triangle A\u2013B\u2013D\u2013A (odd length) makes the graph non-bipartite, so no valid cafe assignment exists.")
    add('Moral of the story, to be able to assign the cafe in the given way, the graph has to be a bipartite graph. With the graph coloring method, it is shown below why the graph is not bipartite.')
    add("As shown below, G\u2013F and B\u2013D are breaking the bipartite condition, so the cafe can't be assigned.")

    add('The Odd Cycles', size=11.5, bold=True, color=DARK, before=8, after=4)
    table([
        ['Cycle', 'Length', 'Odd?'],
        ['A \u2192 B \u2192 D \u2192 A', '3', 'Yes'],
        ['A \u2192 B \u2192 F \u2192 G \u2192 E \u2192 A', '5', 'Yes'],
        ['A \u2192 B \u2192 F \u2192 G \u2192 J \u2192 I \u2192 H \u2192 A', '7', 'Yes'],
    ], widths=[9.0, 2.8, 2.7])
    add('A single odd cycle is enough to break bipartiteness. This graph contains three.')
    image('03_coloring.png', 15.0, '2-colouring attempt. A\u2013B\u2013D\u2013A is an odd cycle (triangle), so the graph is not bipartite.')

    add('3.2 Part B \u2014 Fixing It by Removing a Checkpoint', size=12.5, bold=True, color=BLUE, before=8, after=4)
    add('Yes \u2014 removing checkpoint B makes the assignment possible.', size=11.5, bold=True,
        color=BLUE_DK, fill=CALLOUT_FILL, bar=HEAD_FILL, indent=0.3, after=8)
    add('The cafe assignment is possible with the removal of checkpoint B. The removal of B breaks the odd cycle in the graph. All three odd cycles pass through A and B. Removing either one would have fixed the graph.')
    add('I removed B because of two reasons. Mostly because I wanted to remove it. Secondly, it has a degree of 3 while A has a degree of 5.')

    add('Assignment After Removing B', size=11.5, bold=True, color=DARK, before=8, after=4)
    table([
        ['Espresso Bar', 'Cold Brew Corner'],
        ['A, G, I', 'C, D, E, F, H, J'],
    ], widths=[7.0, 7.0])
    add('Every remaining road joins an Espresso Bar checkpoint to a Cold Brew Corner checkpoint, so no two cafes of the same type share a road.')
    image('04_fixed.png', 15.0, 'After removing B, the graph is bipartite and a valid 2-colouring exists.')

    doc.save(OUT)
    print('saved', OUT)


if __name__ == '__main__':
    build()
